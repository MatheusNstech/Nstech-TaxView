"""Regressões da auditoria de produção: fuso, paginação, import, reabertura, export."""
from __future__ import annotations

import io
from datetime import date, datetime, timezone
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from postgrest.exceptions import APIError

from app.api.routes.obrigacoes import update_obrigacao
from app.core import clock
from app.core.auth import AuthUser
from app.schemas.models import ObrigacaoUpdate
from app.services.csv_import import _import_table_rows
from app.services.db import constraint_errors, fetch_all
from app.services.excel_export import build_obrigacoes_xlsx


class _Query:
    """PostgREST mínimo em memória. Upsert em lote usa a união das colunas
    (colunas ausentes viram NULL), como o postgrest-py faz."""

    def __init__(self, db: dict[str, list[dict]], name: str):
        self.db = db
        self.name = name
        self.op = "select"
        self.payload: list[dict] | dict | None = None
        self.filters: list = []
        self.rng: tuple[int, int] | None = None
        self.conflict: list[str] = []

    def select(self, *_a, **_k):
        return self

    def order(self, *_a, **_k):
        return self

    def limit(self, n):
        self.rng = (0, n - 1)
        return self

    def eq(self, key, value):
        self.filters.append(lambda r: str(r.get(key)) == str(value))
        return self

    def in_(self, key, values):
        allowed = {str(v) for v in values}
        self.filters.append(lambda r: str(r.get(key))[:10] in allowed)
        return self

    def range(self, start, end):
        self.rng = (start, end)
        return self

    def insert(self, payload):
        self.op = "insert"
        self.payload = payload
        return self

    def update(self, payload):
        self.op = "update"
        self.payload = payload
        return self

    def upsert(self, rows, on_conflict: str = "", **_k):
        self.op = "upsert"
        self.payload = rows
        self.conflict = on_conflict.split(",")
        return self

    def execute(self):
        table = self.db.setdefault(self.name, [])
        if self.op == "upsert":
            rows = list(self.payload or [])
            columns = set().union(*(r.keys() for r in rows))
            out = []
            for row in rows:
                full = {c: row.get(c) for c in columns}
                match = next(
                    (
                        t
                        for t in table
                        if all(str(t.get(c)) == str(full.get(c)) for c in self.conflict)
                    ),
                    None,
                )
                if match is None:
                    match = {"id": str(uuid4())}
                    table.append(match)
                match.update(full)
                out.append(dict(match))
            return SimpleNamespace(data=out)
        if self.op == "insert":
            rows = self.payload if isinstance(self.payload, list) else [self.payload]
            out = []
            for row in rows:
                new = {"id": str(uuid4()), **row}
                table.append(new)
                out.append(dict(new))
            return SimpleNamespace(data=out)
        matched = [r for r in table if all(f(r) for f in self.filters)]
        if self.op == "update":
            for row in matched:
                row.update(self.payload or {})
            return SimpleNamespace(data=[dict(r) for r in matched])
        if self.rng is not None:
            matched = matched[self.rng[0] : self.rng[1] + 1]
        return SimpleNamespace(data=[dict(r) for r in matched])


class _FakeDB:
    def __init__(self, **tables: list[dict]):
        self.db = {name: list(rows) for name, rows in tables.items()}

    def table(self, name: str):
        return _Query(self.db, name)


def test_today_br_usa_fuso_de_sao_paulo(monkeypatch):
    class _Fixed(datetime):
        @classmethod
        def now(cls, tz=None):
            # 01:30 UTC do dia 26 = 22:30 do dia 25 em Brasília
            return datetime(2026, 9, 26, 1, 30, tzinfo=timezone.utc).astimezone(tz)

    monkeypatch.setattr(clock, "datetime", _Fixed)
    assert clock.today_br() == date(2026, 9, 25)


def test_fetch_all_pagina_com_builder_novo():
    store = {"t": [{"id": i} for i in range(2500)]}
    builders: list[_Query] = []

    def build():
        q = _Query(store, "t")
        builders.append(q)
        return q

    rows = fetch_all(build, page_size=1000)
    assert len(rows) == 2500
    assert len(builders) == 3
    assert fetch_all(build, page_size=1000, max_rows=1500) == store["t"][:1500]


def test_constraint_errors_vira_409():
    with pytest.raises(HTTPException) as exc:
        with constraint_errors(duplicate="dup"):
            raise APIError({"code": "23505", "message": "duplicate key"})
    assert exc.value.status_code == 409
    assert exc.value.detail == "dup"

    with pytest.raises(APIError):
        with constraint_errors():
            raise APIError({"code": "42P01", "message": "outro erro"})


def test_reimport_preserva_entrega_e_nao_grava_categoria_nula():
    empresa_id, dctf_id, efd_id = str(uuid4()), str(uuid4()), str(uuid4())
    fake = _FakeDB(
        empresas=[{"id": empresa_id, "cnpj": "111"}],
        atividades_modelo=[
            {"id": dctf_id, "nome": "DCTF", "dia_prazo_legal": 15, "dia_prazo_fiscal": 10},
            {"id": efd_id, "nome": "EFD", "dia_prazo_legal": 20, "dia_prazo_fiscal": 15},
        ],
        responsaveis=[],
        obrigacoes=[
            {
                "id": str(uuid4()),
                "empresa_id": empresa_id,
                "atividade_id": dctf_id,
                "competencia": "2026-08-01",
                "status": "ENTREGUE",
                "categoria": "outras",
                "data_entrega": "2026-09-05",
                "recibo_numero": "R-1",
                "observacao": "ok",
                "motivo_atraso": None,
                "responsavel_id": None,
                "prazo_legal": "2026-09-15",
                "prazo_fiscal": "2026-09-10",
            }
        ],
    )
    base = {
        "Tipo": "Obrigação",
        "CNPJ": "111",
        "Razão Social": "Empresa",
        "BU": "X",
        "Competência": "2026-08-01",
        "Status": "",
    }
    rows = [
        {**base, "Atividade": "DCTF", "Categoria": ""},
        {**base, "Atividade": "EFD", "Categoria": "outras"},
    ]

    _import_table_rows(fake, rows, date(2026, 8, 1))  # type: ignore[arg-type]

    obrigacoes = {o["atividade_id"]: o for o in fake.db["obrigacoes"]}
    assert len(obrigacoes) == 2
    dctf = obrigacoes[dctf_id]
    assert dctf["status"] == "ENTREGUE"
    assert dctf["data_entrega"] == "2026-09-05"
    assert dctf["recibo_numero"] == "R-1"
    assert dctf["categoria"] == "outras"
    efd = obrigacoes[efd_id]
    assert efd["status"] == "PENDENTE"
    assert efd["categoria"] == "outras"
    assert all(o["categoria"] for o in obrigacoes.values())


def test_reabrir_obrigacao_entregue_limpa_data_entrega():
    oid = uuid4()
    fake = _FakeDB(
        obrigacoes=[
            {
                "id": str(oid),
                "status": "ENTREGUE",
                "data_entrega": "2026-09-05",
                "prazo_legal": "2099-01-20",
                "prazo_fiscal": "2099-01-15",
            }
        ],
        obrigacao_audit_log=[],
    )
    user = AuthUser(id=str(uuid4()), email="u@x.com", access_token="t", role="user")
    body = ObrigacaoUpdate(status="EM_ANDAMENTO")  # type: ignore[arg-type]

    with patch("app.api.routes.obrigacoes.assert_obrigacao_in_scope"):
        out = update_obrigacao(oid, body, user, fake)  # type: ignore[arg-type]

    assert out["status"] == "EM_ANDAMENTO"
    assert fake.db["obrigacoes"][0]["data_entrega"] is None


def test_export_excel_nao_gera_formula_nem_quebra_com_controle():
    tarefa = {
        "titulo": '=HYPERLINK("http://evil","x")\x07',
        "status": "PENDENTE",
        "prazo": "2026-09-30",
        "empresa": {},
        "responsavel": {},
    }
    content = build_obrigacoes_xlsx([], competencia=None, tarefas=[tarefa])
    ws = load_workbook(io.BytesIO(content)).active
    cell = ws.cell(5, 5)
    assert cell.data_type == "s"
    assert cell.value == '=HYPERLINK("http://evil","x")'
