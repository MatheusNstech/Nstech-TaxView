"""Importa Pendencias.xlsx (fonte) em painel_fiscal_pendencias.

Fallback: Base_Banco_status_empresa.csv

Uso:
  cd backend
  ..\\.venv\\Scripts\\python.exe -m scripts.seed_painel_fiscal
"""

from __future__ import annotations

import csv
import re
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parent
sys.path.insert(0, str(ROOT))

from app.core.auth import get_admin_client  # noqa: E402
from app.core.config import get_settings  # noqa: E402

XLSX_PATH = REPO / "Pendencias.xlsx"
CSV_PATH = REPO / "Base_Banco_status_empresa.csv"


def parse_money(raw: object) -> Decimal | None:
    """Aceita número Excel ou texto BR (1.030,77)."""
    if raw is None or raw == "":
        return None
    if isinstance(raw, (int, float)):
        return Decimal(str(raw)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    text = str(raw).strip()
    if not text or text in {"-", "—"}:
        return None
    text = text.replace("R$", "").replace(" ", "").strip()
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except InvalidOperation:
        return None


def parse_int(raw: object) -> int | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, int):
        return raw
    if isinstance(raw, float):
        return int(raw)
    text = str(raw).strip()
    if not text:
        return None
    try:
        return int(float(text.replace(",", ".")))
    except ValueError:
        return None


def parse_date(raw: object) -> str | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    text = str(raw).strip()
    if not text:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def clean_cnpj(raw: object) -> str:
    return re.sub(r"\D", "", str(raw or "").strip())


def normalize_orgao(raw: object) -> str:
    key = str(raw or "").strip().upper()
    if "CADIN" in key:
        return "CADIN"
    if "PGFN" in key:
        return "PGFN"
    if key == "RFB":
        return "RFB"
    if "SEM PEND" in key:
        return "Sem Pendência"
    return str(raw or "").strip()


def _get(row: dict, *names: str) -> object:
    for name in names:
        if name in row and row[name] is not None:
            return row[name]
    # encoding-tolerant match
    lower = {str(k).casefold(): k for k in row}
    for name in names:
        key = lower.get(name.casefold())
        if key is not None:
            return row[key]
    for k, v in row.items():
        kn = str(k)
        for name in names:
            if name.casefold() in kn.casefold():
                return v
    return None


def row_to_payload(row: dict) -> dict:
    return {
        "empresa": str(_get(row, "Empresa") or "").strip() or "—",
        "razao_social": str(_get(row, "Razão Social", "Razao Social") or "").strip(),
        "situacao_cnpj": str(_get(row, "Situação CNPJ", "Situacao CNPJ") or "").strip(),
        "cnpj": clean_cnpj(_get(row, "CNPJ")),
        "cidade_iss": str(_get(row, "Cidade (ISS)") or "").strip(),
        "uf": str(_get(row, "UF") or "").strip(),
        "orgao": normalize_orgao(_get(row, "Orgão", "Orgao", "Órgão")),
        "sucedida": str(_get(row, "Sucedida") or "").strip(),
        "data_inscricao": parse_date(_get(row, "Data da Inscrição", "Data da Inscricao")),
        "cnpj_sucedida": clean_cnpj(_get(row, "CNPJ_Sucedida")),
        "empresa_sucedida": str(_get(row, "Empresa Sucedida") or "").strip(),
        "natureza": str(_get(row, "Natureza") or "").strip(),
        "fase": str(_get(row, "Fase") or "").strip(),
        "tipo": str(_get(row, "Tipo") or "").strip(),
        "situacao": str(_get(row, "Situação", "Situacao") or "").strip(),
        "codigo": str(_get(row, "Cód", "Cod", "Cód.") or "").strip(),
        "mes": parse_int(_get(row, "Mês", "Mes")),
        "ano": parse_int(_get(row, "Ano")),
        "periodo_apuracao": str(
            _get(row, "Período de Apuração", "Periodo de Apuracao") or ""
        ).strip(),
        "vencimento": parse_date(_get(row, "Vencimento")),
        "principal": float(v) if (v := parse_money(_get(row, "Principal"))) is not None else None,
        "multa": float(v) if (v := parse_money(_get(row, "Multa"))) is not None else None,
        "juros": float(v) if (v := parse_money(_get(row, "Juros"))) is not None else None,
        "total": float(v) if (v := parse_money(_get(row, "Total"))) is not None else None,
        "motivo": str(_get(row, "Motivo") or "").strip(),
        "numero_processo": str(
            _get(row, "Nº Processo ", "Nº Processo", "No Processo") or ""
        ).strip(),
        "cnd": str(_get(row, "CND") or "").strip(),
        "validade_cnd": parse_date(_get(row, "Validade")),
        "status_cnd": str(_get(row, "Status CND") or "").strip(),
        "nota_01": str(_get(row, "Nota_01") or "").strip(),
        "nota_02": str(_get(row, "Nota_02") or "").strip(),
    }


def load_rows_from_xlsx(path: Path) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        raw = {headers[i]: values[i] for i in range(len(headers))}
        if not str(raw.get("Empresa") or "").strip():
            continue
        rows.append(row_to_payload(raw))
    return rows


def load_rows_from_csv(path: Path) -> list[dict]:
    with path.open(encoding="latin-1", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=";")
        return [row_to_payload(r) for r in reader if (r.get("Empresa") or "").strip()]


def main() -> int:
    settings = get_settings()
    if not settings.supabase_service_role_key or settings.supabase_service_role_key.startswith(
        "your-"
    ):
        print("Configure SUPABASE_SERVICE_ROLE_KEY")
        return 1

    if XLSX_PATH.exists():
        print(f"Fonte: {XLSX_PATH.name}")
        rows = load_rows_from_xlsx(XLSX_PATH)
    elif CSV_PATH.exists():
        print(f"Fonte (fallback): {CSV_PATH.name}")
        rows = load_rows_from_csv(CSV_PATH)
    else:
        print("Nem Pendencias.xlsx nem Base_Banco_status_empresa.csv encontrados")
        return 1

    admin = get_admin_client(settings)
    admin.table("painel_fiscal_pendencias").delete().neq("empresa", "").execute()

    batch = 50
    inserted = 0
    for i in range(0, len(rows), batch):
        chunk = rows[i : i + batch]
        admin.table("painel_fiscal_pendencias").insert(chunk).execute()
        inserted += len(chunk)

    # sanity vs painel Excel
    by: dict[str, float] = {}
    for r in rows:
        org = r["orgao"] or "—"
        by[org] = by.get(org, 0.0) + float(r["total"] or 0)
    print(f"Importadas {inserted} linhas em painel_fiscal_pendencias")
    for org in ("CADIN", "PGFN", "RFB", "Sem Pendência"):
        print(f"  {org}: R$ {by.get(org, 0.0):,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
