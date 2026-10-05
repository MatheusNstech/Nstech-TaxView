"""Gera os layouts de importação com linhas de exemplo.

- Cronograma (obrigações e tarefas): cabeçalhos de IMPORT_HEADERS, entram pela tela Importação.
- Painel fiscal e PER/DCOMP (área do Danilo): cabeçalhos lidos por seed_painel_fiscal e
  seed_perdcomp, entram rodando esses scripts.

Uso:
  cd backend
  ..\\.venv\\Scripts\\python.exe -m scripts.gerar_layouts_importacao ["pasta de saída"]
"""

from __future__ import annotations

import sys
from dataclasses import dataclass, field
from datetime import date, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.services.csv_import import IMPORT_HEADERS  # noqa: E402

DEFAULT_OUT = (
    Path.home()
    / "OneDrive - NSTECH GR LTDA"
    / "Documentos"
    / "Cronograma"
    / "Exemplos de Importação"
)

COMPETENCIA = date(2026, 9, 1)
PRAZO_LEGAL = date(2026, 10, 20)
PRAZO_FISCAL = date(2026, 10, 15)
MAX_ROWS = 1000

FILL_OBRIGATORIO = PatternFill("solid", fgColor="FF3D03")
FILL_OPCIONAL = PatternFill("solid", fgColor="595959")
FILL_DEPENDE = PatternFill("solid", fgColor="1F4E79")
FILL_NAO_APLICA = PatternFill("solid", fgColor="D9D9D9")
FILL_EXEMPLO = PatternFill("solid", fgColor="FFF2CC")
FONT_HEADER = Font(bold=True, color="FFFFFF")
FONT_NAO_APLICA = Font(bold=True, color="808080")
FONT_EXEMPLO = Font(italic=True, color="7F6000")

OBRIGATORIO = "Sim"
OPCIONAL = "Não"
NAO_APLICA = "Não se aplica"
DEPENDE = "Depende do tipo"

STATUS_LIST = ["PENDENTE", "EM_ANDAMENTO", "EM_REVISAO", "ENTREGUE", "ATRASADO"]


@dataclass
class Coluna:
    nome: str
    obrigatorio: str
    formato: str
    tipo: str = "texto"  # texto | data | hora | numero | inteiro
    lista: list[str] | None = None
    largura: int | None = None
    por_tipo: tuple[str, str] | None = None  # (Obrigação, Tarefa)


@dataclass
class Aba:
    titulo: str
    colunas: list[Coluna]
    exemplos: list[dict[str, object]]
    observacoes: list[str] = field(default_factory=list)


@dataclass
class Layout:
    arquivo: str
    titulo: str
    como_carregar: list[str]
    regras: list[str]
    abas: list[Aba]


def _number_format(tipo: str) -> str | None:
    return {
        "texto": "@",
        "data": "DD/MM/YYYY",
        "hora": "HH:MM",
        "numero": "#,##0.00",
        "inteiro": "0",
    }.get(tipo)


def _write_data_sheet(ws: Worksheet, aba: Aba) -> None:
    ws.append([c.nome for c in aba.colunas])
    for exemplo in aba.exemplos:
        ws.append([exemplo.get(c.nome) for c in aba.colunas])

    for idx, col in enumerate(aba.colunas, start=1):
        letter = get_column_letter(idx)
        header = ws.cell(1, idx)
        if col.obrigatorio == NAO_APLICA:
            header.fill, header.font = FILL_NAO_APLICA, FONT_NAO_APLICA
        elif col.obrigatorio == OBRIGATORIO:
            header.fill, header.font = FILL_OBRIGATORIO, FONT_HEADER
        elif col.obrigatorio == DEPENDE:
            header.fill, header.font = FILL_DEPENDE, FONT_HEADER
        else:
            header.fill, header.font = FILL_OPCIONAL, FONT_HEADER
        header.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[letter].width = col.largura or max(16, len(col.nome) + 4)

        fmt = _number_format(col.tipo)
        if fmt:
            for row_idx in range(2, MAX_ROWS + 1):
                ws.cell(row_idx, idx).number_format = fmt

        if col.lista:
            dv = DataValidation(
                type="list",
                formula1='"' + ",".join(col.lista) + '"',
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="Valor inválido",
                error="Escolha um valor da lista: " + ", ".join(col.lista),
            )
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{MAX_ROWS}")

    for row_idx in range(2, len(aba.exemplos) + 2):
        for idx in range(1, len(aba.colunas) + 1):
            cell = ws.cell(row_idx, idx)
            cell.fill, cell.font = FILL_EXEMPLO, FONT_EXEMPLO

    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(aba.colunas))}1"


def _write_instructions(ws: Worksheet, layout: Layout) -> None:
    ws["A1"] = layout.titulo
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Competência de referência: {COMPETENCIA:%m/%Y}"
    ws["A2"].font = Font(italic=True, color="595959")

    row = 4
    ws.cell(row, 1, "Como carregar").font = Font(bold=True, size=12)
    row += 1
    for linha in layout.como_carregar:
        ws.cell(row, 1, f"- {linha}")
        row += 1

    row += 1
    ws.cell(row, 1, "Regras importantes").font = Font(bold=True, size=12)
    row += 1
    for linha in [
        "As linhas em amarelo são exemplos: apague-as antes de carregar.",
        "Cabeçalho laranja = obrigatório; azul = obrigatório só para um dos tipos (veja a tabela abaixo); "
        "cinza escuro = opcional; cinza claro = não se aplica (deixe em branco). "
        "Não renomeie nem mude a ordem das colunas.",
        *layout.regras,
    ]:
        ws.cell(row, 1, f"- {linha}")
        row += 1

    for aba in layout.abas:
        row += 1
        ws.cell(row, 1, f"Colunas da aba \"{aba.titulo}\"").font = Font(bold=True, size=12)
        row += 1
        for linha in aba.observacoes:
            ws.cell(row, 1, f"- {linha}")
            row += 1
        por_tipo = any(col.por_tipo for col in aba.colunas)
        flags = ["Obrigação", "Tarefa"] if por_tipo else ["Obrigatório"]
        for idx, titulo in enumerate(["Coluna", *flags, "Formato / valores aceitos", "Exemplo"], 1):
            cell = ws.cell(row, idx, titulo)
            cell.fill, cell.font = FILL_OPCIONAL, FONT_HEADER
        row += 1
        for col in aba.colunas:
            valor = next((e[col.nome] for e in aba.exemplos if e.get(col.nome) not in (None, "")), None)
            if isinstance(valor, date):
                valor = valor.strftime("%d/%m/%Y")
            elif isinstance(valor, time):
                valor = valor.strftime("%H:%M")
            valores = [col.nome, *(col.por_tipo or (col.obrigatorio,)), col.formato,
                       "" if valor is None else str(valor)]
            for idx, texto in enumerate(valores, 1):
                ws.cell(row, idx, texto)
            ws.cell(row, 1).font = Font(bold=True)
            ws.cell(row, len(valores) - 1).alignment = Alignment(wrap_text=True, vertical="top")
            row += 1

    por_tipo = any(col.por_tipo for aba in layout.abas for col in aba.colunas)
    larguras = [34, 16, 16, 80, 40] if por_tipo else [34, 16, 90, 40]
    for idx, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(idx)].width = largura


def save_layout(layout: Layout, out_dir: Path) -> Path:
    wb = Workbook()
    first = True
    for aba in layout.abas:
        ws = wb.active if first else wb.create_sheet()
        ws.title = aba.titulo
        _write_data_sheet(ws, aba)
        first = False
    _write_instructions(wb.create_sheet("Instruções"), layout)
    path = out_dir / layout.arquivo
    wb.save(path)
    return path


CNPJ_FMT = "Texto no formato 00.000.000/0000-00, idêntico ao cadastro de empresas."

IMPORT_FORMATOS: dict[str, tuple[str, str]] = {
    "Tipo": ("texto", "Obrigação ou Tarefa."),
    "CNPJ": ("texto", CNPJ_FMT),
    "Razão Social": ("texto", "Usado só se a empresa ainda não existir."),
    "BU": ("texto", "Usado só se a empresa ainda não existir."),
    "Apuração": ("texto", "VERDADEIRO ou FALSO. Usado só se a atividade ainda não existir."),
    "Atividade": ("texto", "Nome exato da atividade cadastrada (ex.: Apuração ISS Prestados)."),
    "Título": ("texto", "Até 200 caracteres."),
    "Responsável": ("texto", "Nome exato do responsável cadastrado; um nome diferente cria um responsável novo."),
    "E-mail responsável": ("texto", "Usado só se o responsável ainda não existir."),
    "Solicitante": ("texto", "Quem pediu a tarefa (2 a 120 caracteres)."),
    "Descrição": ("texto", "Texto livre."),
    "Categoria": ("texto", "fechamento ou outras."),
    "Recorrência": ("texto", "mensal, trimestral ou anual. Usado só se a atividade ainda não existir."),
    "Dia prazo legal": ("inteiro", "1 a 31. Usado só se a atividade ainda não existir (padrão 20)."),
    "Dia prazo fiscal": ("inteiro", "1 a 31. Usado só se a atividade ainda não existir (padrão 15)."),
    "Competência": ("data", "Data dd/mm/aaaa, sempre dia 01 (01/09/2026)."),
    "Prazo legal": ("data", "Data dd/mm/aaaa. Vazio = dia prazo legal do mês seguinte."),
    "Prazo fiscal": ("data", "Data dd/mm/aaaa. Vazio = dia prazo fiscal do mês seguinte."),
    "Hora início": ("hora", "HH:MM."),
    "Hora fim": ("hora", "HH:MM, maior que Hora início."),
    "Data entrega": ("data", "Data dd/mm/aaaa. Preencha quando o Status for ENTREGUE."),
    "Número recibo": ("texto", "Número do recibo de entrega."),
    "Observação": ("texto", "Texto livre."),
    "Motivo atraso": ("texto", "Obrigatório quando entregue após o prazo ou ATRASADO (até 2000 caracteres)."),
    "Status": ("texto", ", ".join(STATUS_LIST) + "."),
}

IMPORT_LISTAS: dict[str, list[str]] = {
    "Tipo": ["Obrigação", "Tarefa"],
    "Apuração": ["VERDADEIRO", "FALSO"],
    "Categoria": ["fechamento", "outras"],
    "Recorrência": ["mensal", "trimestral", "anual"],
    "Status": STATUS_LIST,
}

IMPORT_LARGURAS = {"Razão Social": 32, "Atividade": 30, "Título": 32, "Descrição": 32,
                   "Observação": 32, "Motivo atraso": 40, "CNPJ": 22, "E-mail responsável": 30}


OBRIGATORIAS_OBRIGACAO = {"Tipo", "CNPJ", "Atividade", "Responsável", "Competência", "Status"}
NAO_APLICA_OBRIGACAO = {"Título", "Solicitante", "Descrição", "Hora início", "Hora fim"}
OBRIGATORIAS_TAREFA = {"Tipo", "Título", "Responsável", "Solicitante", "Prazo fiscal",
                       "Hora início", "Hora fim", "Status"}
NAO_APLICA_TAREFA = {"Apuração", "Atividade", "Recorrência", "Dia prazo legal", "Dia prazo fiscal",
                     "Data entrega", "Número recibo", "Observação"}


def _flag(nome: str, obrigatorias: set[str], nao_aplica: set[str]) -> str:
    if nome in obrigatorias:
        return OBRIGATORIO
    if nome in nao_aplica:
        return NAO_APLICA
    return OPCIONAL


def _cronograma_colunas() -> list[Coluna]:
    colunas = []
    for nome in IMPORT_HEADERS:
        tipo, formato = IMPORT_FORMATOS[nome]
        por_tipo = (
            _flag(nome, OBRIGATORIAS_OBRIGACAO, NAO_APLICA_OBRIGACAO),
            _flag(nome, OBRIGATORIAS_TAREFA, NAO_APLICA_TAREFA),
        )
        if por_tipo == (OBRIGATORIO, OBRIGATORIO):
            geral = OBRIGATORIO
        elif OBRIGATORIO in por_tipo:
            geral = DEPENDE
        else:
            geral = OPCIONAL
        colunas.append(Coluna(nome, geral, formato, tipo=tipo, lista=IMPORT_LISTAS.get(nome),
                              largura=IMPORT_LARGURAS.get(nome), por_tipo=por_tipo))
    return colunas


def layout_cronograma() -> Layout:
    empresa = {"CNPJ": "00.000.000/0001-00", "Razão Social": "Empresa Exemplo LTDA", "BU": "BU Exemplo",
               "Responsável": "Nome do Responsável", "E-mail responsável": "responsavel@nstech.com.br",
               "Competência": COMPETENCIA}
    obrigacao = {**empresa, "Tipo": "Obrigação", "Apuração": "VERDADEIRO", "Categoria": "fechamento",
                 "Recorrência": "mensal", "Dia prazo legal": 20, "Dia prazo fiscal": 15,
                 "Prazo legal": PRAZO_LEGAL, "Prazo fiscal": PRAZO_FISCAL}
    tarefa = {**empresa, "Tipo": "Tarefa"}
    exemplos = [
        {**obrigacao, "Atividade": "Apuração ISS Prestados", "Status": "PENDENTE"},
        {**obrigacao, "Atividade": "EFD Contribuições", "Status": "ENTREGUE", "Data entrega": date(2026, 10, 10),
         "Número recibo": "12.34.56.78.90-12", "Observação": "Entregue sem pendências."},
        {**obrigacao, "Atividade": "REINF", "Status": "ATRASADO",
         "Motivo atraso": "Aguardando o cliente enviar as notas com retenção para fechar a apuração."},
        {**tarefa, "Título": "Conferir guias do cliente", "Solicitante": "Maria Solicitante",
         "Descrição": "Follow-up pedido pelo cliente", "Categoria": "outras",
         "Prazo fiscal": date(2026, 9, 28), "Hora início": time(9, 0), "Hora fim": time(12, 0),
         "Status": "PENDENTE"},
        {**tarefa, "CNPJ": None, "Razão Social": None, "BU": None,
         "Título": "Reunião de fechamento ISS", "Solicitante": "Coordenação Fiscal",
         "Categoria": "fechamento", "Prazo fiscal": date(2026, 10, 1),
         "Hora início": time(14, 0), "Hora fim": time(15, 30), "Status": "EM_ANDAMENTO"},
    ]
    aba = Aba(
        "Importação",
        _cronograma_colunas(),
        exemplos,
        observacoes=[
            "Obrigação: uma linha por empresa + atividade da competência 09/2026 (entrega fiscal recorrente).",
            "Tarefa: uma linha por trabalho avulso pedido por alguém, com prazo e horário. "
            "CNPJ é opcional (tarefa interna pode ficar sem empresa).",
        ],
    )
    return Layout(
        "01_Cronograma_Setembro.xlsx",
        "Cronograma (obrigações e tarefas) - competência 09/2026",
        [
            "Tela Importação (admin): escolha a competência 09/2026 e envie o arquivo .xlsx (até 10 MB).",
            "A aba precisa se chamar \"Importação\" (já está assim).",
            "Obrigações e tarefas vão na mesma aba: a coluna Tipo diz o que é cada linha.",
        ],
        [
            "Obrigação: a importação atualiza a obrigação existente da mesma empresa + atividade + competência "
            "(não duplica): status, prazos e responsável são sobrescritos pelo que estiver na planilha.",
            "Obrigação: Responsável vazio APAGA o responsável já atribuído; preencha sempre.",
            "Obrigação: CNPJ e Atividade devem ser idênticos ao cadastro; um nome diferente cria uma atividade "
            "nova. Linhas sem CNPJ ou sem Atividade são ignoradas.",
            "Obrigação ENTREGUE: preencha Data entrega (e Número recibo, se houver).",
            "Tarefa: a importação só inclui tarefas novas; não atualiza tarefas já cadastradas. "
            "Tarefa com mesmo título + responsável + prazo + hora início + empresa é ignorada.",
            "Tarefa: linhas sem Título, Responsável, Solicitante, Prazo fiscal, Hora início ou Hora fim são "
            "ignoradas. Hora fim precisa ser maior que Hora início.",
            "Responsável (obrigação ou tarefa): use o nome idêntico ao cadastro (Felipe, Flávia, Glaucia, "
            "Solange, Viviane...); um nome diferente cria um responsável novo.",
            "Entregue após o prazo ou ATRASADO (obrigação ou tarefa): preencha Motivo atraso.",
        ],
        [aba],
    )


def layout_painel_fiscal() -> Layout:
    moeda = "Valor em R$ (número ou texto 1.234,56). \"-\" = vazio."
    colunas = [
        Coluna("Empresa", OBRIGATORIO, "Nome curto da empresa. Linha sem Empresa é ignorada.", largura=22),
        Coluna("Razão Social", OPCIONAL, "Texto.", largura=32),
        Coluna("Situação CNPJ", OPCIONAL, "Ex.: Ativa, Baixada.", lista=["Ativa", "Baixada", "Inapta", "Suspensa"]),
        Coluna("CNPJ", OBRIGATORIO, "CNPJ com ou sem pontuação.", largura=22),
        Coluna("Cidade (ISS)", OPCIONAL, "Município do ISS."),
        Coluna("UF", OPCIONAL, "Sigla do estado.", largura=8),
        Coluna("Orgão", OBRIGATORIO, "CADIN, PGFN, RFB ou Sem Pendência.",
               lista=["CADIN", "PGFN", "RFB", "Sem Pendência"]),
        Coluna("Sucedida", OPCIONAL, "Sim ou Não.", lista=["Sim", "Não"]),
        Coluna("Data da Inscrição", OPCIONAL, "Data dd/mm/aaaa.", tipo="data"),
        Coluna("CNPJ_Sucedida", OPCIONAL, "CNPJ da empresa sucedida.", largura=22),
        Coluna("Empresa Sucedida", OPCIONAL, "Texto."),
        Coluna("Natureza", OPCIONAL, "Ex.: Tributário, Previdenciário."),
        Coluna("Fase", OPCIONAL, "Fase da pendência (ex.: Sem Pendência, Inscrita, Parcelada)."),
        Coluna("Tipo", OPCIONAL, "Tipo da pendência (ex.: Débito, Omissão, Não se aplica)."),
        Coluna("Situação", OPCIONAL, "Situação da pendência."),
        Coluna("Cód", OPCIONAL, "Código da receita/débito.", largura=10),
        Coluna("Mês", OPCIONAL, "Número de 1 a 12.", tipo="inteiro", largura=8),
        Coluna("Ano", OPCIONAL, "Ano com 4 dígitos.", tipo="inteiro", largura=8),
        Coluna("Período de Apuração", OPCIONAL, "Texto (ex.: 09/2026)."),
        Coluna("Vencimento", OPCIONAL, "Data dd/mm/aaaa.", tipo="data"),
        Coluna("Principal", OPCIONAL, moeda, tipo="numero"),
        Coluna("Multa", OPCIONAL, moeda, tipo="numero"),
        Coluna("Juros", OPCIONAL, moeda, tipo="numero"),
        Coluna("Total", OPCIONAL, moeda, tipo="numero"),
        Coluna("Motivo", OPCIONAL, "Texto livre.", largura=32),
        Coluna("Nº Processo ", OPCIONAL, "Número do processo (o nome da coluna termina com espaço, não remova)."),
        Coluna("CND", OPCIONAL, "Ex.: Certidão Negativa, Positiva com Efeitos de Negativa.", largura=28),
        Coluna("Validade", OPCIONAL, "Validade da CND, data dd/mm/aaaa.", tipo="data"),
        Coluna("Status CND", OPCIONAL, "Ex.: Válida, Vencida.", lista=["Válida", "Vencida"]),
        Coluna("Nota_01", OPCIONAL, "Anotação livre.", largura=28),
        Coluna("Nota_02", OPCIONAL, "Anotação livre.", largura=28),
    ]
    base = {"Empresa": "EXEMPLO", "Razão Social": "Empresa Exemplo LTDA", "Situação CNPJ": "Ativa",
            "CNPJ": "00.000.000/0001-00", "Cidade (ISS)": "São Paulo", "UF": "SP", "Sucedida": "Não",
            "Natureza": "Tributário"}
    exemplos = [
        {**base, "Orgão": "Sem Pendência", "Fase": "Sem Pendência", "Tipo": "Não se aplica",
         "Situação": "Não se aplica", "Mês": 9, "Ano": 2026, "CND": "Certidão Negativa",
         "Validade": date(2027, 2, 8), "Status CND": "Válida"},
        {**base, "Orgão": "PGFN", "Fase": "Inscrita", "Tipo": "Débito", "Situação": "Em cobrança",
         "Cód": "8109", "Mês": 9, "Ano": 2026, "Período de Apuração": "08/2026",
         "Vencimento": date(2026, 9, 25), "Principal": 1000.00, "Multa": 200.00, "Juros": 30.77,
         "Total": 1230.77, "Motivo": "Divergência DCTFWeb x DARF", "Nº Processo ": "10880.000000/2026-00",
         "CND": "Positiva com Efeitos de Negativa", "Validade": date(2026, 12, 31), "Status CND": "Válida"},
    ]
    aba = Aba("Pendencias", colunas, exemplos,
              observacoes=["Uma linha por pendência; empresa sem pendência fica com uma linha \"Sem Pendência\"."])
    return Layout(
        "05_Painel_Fiscal_Pendencias.xlsx",
        "Painel fiscal: pendências e CND (responsável: Danilo)",
        [
            "Danilo preenche e envia o arquivo para um admin.",
            "O admin salva como Pendencias.xlsx na raiz do projeto e roda: "
            "cd backend; ..\\.venv\\Scripts\\python.exe -m scripts.seed_painel_fiscal",
            "Ajustes pontuais depois da carga: tela Painel Fiscal (o Danilo tem permissão de edição).",
        ],
        [
            "ATENÇÃO: a carga APAGA toda a base do painel fiscal e recarrega com esta planilha. "
            "Ela precisa conter todas as pendências de todas as empresas, não só as de setembro.",
            "A aba de dados precisa ser a primeira do arquivo.",
        ],
        [aba],
    )


def layout_perdcomp() -> Layout:
    colunas = [
        Coluna("PER/DCOMP", OBRIGATORIO, "Número do PER/DCOMP (chave única). Linha sem número é ignorada.",
               largura=34),
        Coluna("Processo", OPCIONAL, "Número do processo administrativo.", largura=24),
        Coluna("Empresa", OBRIGATORIO, "Nome da empresa.", largura=24),
        Coluna("Tributo/Crédito", OBRIGATORIO, "Ex.: Saldo negativo IRPJ, PIS/COFINS.", largura=24),
        Coluna("Período", OPCIONAL, "Período do crédito (texto, ex.: 3º trim/2025)."),
        Coluna("Valor do Pedido", OBRIGATORIO, "Valor em R$ (número ou texto 1.234,56).", tipo="numero"),
        Coluna("Status", OPCIONAL, "Situação atual (ex.: Em análise, Deferido, Indeferido, Intimação/Pendência).",
               lista=["Em análise", "Deferido", "Deferido parcialmente", "Indeferido", "Intimação/Pendência"]),
        Coluna("Observações", OPCIONAL, "Texto livre.", largura=32),
        Coluna("Prazo para Cumprimento", OPCIONAL, "Texto (ex.: 30 dias)."),
        Coluna("Data-base / Ciência", OPCIONAL, "Texto livre (ex.: Intimação lida em 03/09/2026).", largura=32),
        Coluna("Data Limite", OPCIONAL, "Data dd/mm/aaaa (usada no alerta de prazo).", tipo="data"),
        Coluna("Providência / Observação do Prazo", OPCIONAL, "Texto livre.", largura=36),
    ]
    exemplos = [
        {"PER/DCOMP": "12345.12345.120926.1.3.02-1234", "Processo": "10880.000000/2026-00",
         "Empresa": "Empresa Exemplo LTDA", "Tributo/Crédito": "Saldo negativo IRPJ", "Período": "2025",
         "Valor do Pedido": 250000.00, "Status": "Intimação/Pendência",
         "Observações": "Intimação para apresentar documentos.", "Prazo para Cumprimento": "30 dias",
         "Data-base / Ciência": "Intimação lida em 03/09/2026", "Data Limite": date(2026, 10, 3),
         "Providência / Observação do Prazo": "Separar razão contábil e informes de rendimento."},
        {"PER/DCOMP": "54321.54321.010926.1.1.01-4321", "Empresa": "Empresa Exemplo LTDA",
         "Tributo/Crédito": "PIS/COFINS", "Período": "08/2026", "Valor do Pedido": 18500.00,
         "Status": "Em análise"},
    ]
    aba = Aba("Controle PERDCOMP", colunas, exemplos,
              observacoes=["Uma linha por PER/DCOMP."])
    return Layout(
        "06_PERDCOMP.xlsx",
        "Controle de PER/DCOMP (responsável: Danilo)",
        [
            "Danilo preenche e envia o arquivo para um admin.",
            "O admin roda: cd backend; ..\\.venv\\Scripts\\python.exe -m scripts.seed_perdcomp "
            "\"<caminho do arquivo>\"",
            "Ajustes pontuais depois da carga: tela PER/DCOMP (o Danilo tem permissão de edição).",
        ],
        [
            "A carga atualiza pelo número do PER/DCOMP: reenviar o mesmo número sobrescreve a linha (não duplica).",
            "Pode conter só os processos novos ou alterados; os que não estiverem na planilha continuam como estão.",
            "A aba de dados precisa ser a primeira do arquivo.",
        ],
        [aba],
    )


LAYOUTS = [layout_cronograma, layout_painel_fiscal, layout_perdcomp]


def main(argv: list[str]) -> int:
    out_dir = Path(argv[1]).expanduser() if len(argv) > 1 else DEFAULT_OUT
    out_dir.mkdir(parents=True, exist_ok=True)
    for build in LAYOUTS:
        path = save_layout(build(), out_dir)
        print(f"Gerado: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
