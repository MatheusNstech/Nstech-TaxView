from __future__ import annotations

import io
from calendar import monthrange
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

BRAND = RGBColor(0xFF, 0x3C, 0x02)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xFF, 0xF7, 0xF5)
LINE = RGBColor(0xF3, 0xE8, 0xE4)
DARK = RGBColor(0x1A, 0x12, 0x10)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
LINES_PER_SLIDE = 10

STATUS_LABELS = {
    "PENDENTE": "Pendente",
    "EM_ANDAMENTO": "Em andamento",
    "EM_REVISAO": "Em revisão",
    "ENTREGUE": "Entregue",
    "ATRASADO": "Atrasado",
}
STATUS_COLORS = {
    "PENDENTE": RGBColor(0x64, 0x74, 0x8B),
    "EM_ANDAMENTO": RGBColor(0x02, 0x84, 0xC7),
    "EM_REVISAO": RGBColor(0xD9, 0x77, 0x06),
    "ENTREGUE": RGBColor(0x05, 0x96, 0x69),
    "ATRASADO": RGBColor(0xE1, 0x1D, 0x48),
}
WEEKDAYS = ("segunda", "terça", "quarta", "quinta", "sexta", "sábado", "domingo")
MONTHS = (
    "janeiro",
    "fevereiro",
    "março",
    "abril",
    "maio",
    "junho",
    "julho",
    "agosto",
    "setembro",
    "outubro",
    "novembro",
    "dezembro",
)


@dataclass
class CalendarLine:
    title: str
    horario: str
    status: str
    empresas: int
    nomes: tuple[str, ...] = ()


@dataclass
class DayBucket:
    day: date
    offset: int
    lines: list[CalendarLine] = field(default_factory=list)


@dataclass
class PostMortemItem:
    title: str
    status: str
    prazo: date
    empresas: int
    detalhe: str


def market_call_filename(competencia: date | None) -> str:
    if competencia is None:
        return "Market Call BR.pptx"
    return (
        f"{competencia.month:02d}.{competencia.year} - "
        f"Market Call BR - P{competencia.month}.pptx"
    )


def _as_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value)[:10]
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def _hora(value: Any) -> str:
    if value is None or value == "":
        return ""
    text = str(value).strip()
    return text[:5]


def _horario(row: dict[str, Any]) -> str:
    start = _hora(row.get("hora_inicio"))
    end = _hora(row.get("hora_fim"))
    if start and end:
        return f"{start}–{end}"
    return start or end


def _empresa_key(row: dict[str, Any]) -> str:
    emp = row.get("empresa") or {}
    return str(emp.get("id") or emp.get("cnpj") or emp.get("razao_social") or "—")


def _empresa_nome(row: dict[str, Any]) -> str:
    emp = row.get("empresa") or {}
    return str(emp.get("razao_social") or "").strip()


def _month_end(competencia: date) -> date:
    return date(
        competencia.year,
        competencia.month,
        monthrange(competencia.year, competencia.month)[1],
    )


def _business_offset(anchor: date, prazo: date) -> int:
    if prazo == anchor:
        return 0
    step = 1 if prazo > anchor else -1
    current = anchor
    count = 0
    while current != prazo:
        current += timedelta(days=step)
        if current.weekday() < 5:
            count += step
    return count


def _bd_label(offset: int) -> str:
    if offset == 0:
        return "BD"
    sign = "+" if offset > 0 else "-"
    return f"BD {sign}{abs(offset)}"


def _period_label(competencia: date | None) -> str:
    if competencia is None:
        return "Fechamento"
    return f"Fechamento P{competencia.month}/{competencia.year % 100:02d}"


def _resolve_anchor(competencia: date | None, rows: list[dict[str, Any]]) -> date:
    if competencia is not None:
        return _month_end(competencia)
    dates = [_as_date(row.get("competencia")) for row in rows]
    known = [item for item in dates if item is not None]
    if known:
        return _month_end(known[0])
    return _month_end(date.today())


def _status(row: dict[str, Any]) -> str:
    return str(row.get("status") or "PENDENTE")


def _obrigacao_prazo(row: dict[str, Any]) -> date | None:
    return _as_date(row.get("prazo_fiscal")) or _as_date(row.get("prazo_legal"))


def _title_obrigacao(row: dict[str, Any]) -> str:
    atividade = row.get("atividade") or {}
    return str(atividade.get("nome") or "Obrigação").strip() or "Obrigação"


def _title_tarefa(row: dict[str, Any]) -> str:
    return str(row.get("titulo") or "Tarefa").strip() or "Tarefa"


def _late_delivery(row: dict[str, Any], prazo: date, kind: str) -> date | None:
    if kind == "tarefa":
        delivered = _as_date(row.get("entregue_em"))
    else:
        delivered = _as_date(row.get("data_entrega"))
    if delivered and delivered > prazo:
        return delivered
    return None


def _collect_items(
    obrigacoes: list[dict[str, Any]],
    tarefas: list[dict[str, Any]],
) -> list[tuple[str, dict[str, Any], date | None, str, str]]:
    items: list[tuple[str, dict[str, Any], date | None, str, str]] = []
    for row in obrigacoes:
        items.append(("obrigacao", row, _obrigacao_prazo(row), _title_obrigacao(row), ""))
    for row in tarefas:
        items.append(("tarefa", row, _as_date(row.get("prazo")), _title_tarefa(row), _horario(row)))
    return items


def _build_days(
    items: list[tuple[str, dict[str, Any], date | None, str, str]],
    anchor: date,
) -> list[DayBucket]:
    grouped: dict[date, dict[tuple[str, str, str], set[tuple[str, str]]]] = defaultdict(
        lambda: defaultdict(set)
    )
    for _kind, row, prazo, title, horario in items:
        if prazo is None:
            continue
        status = _status(row)
        grouped[prazo][(title, horario, status)].add((_empresa_key(row), _empresa_nome(row)))

    days: list[DayBucket] = []
    for day in sorted(grouped):
        lines = []
        for (title, horario, status), empresas in grouped[day].items():
            nomes = tuple(sorted({nome for _key, nome in empresas if nome}))
            lines.append(
                CalendarLine(
                    title=title,
                    horario=horario,
                    status=status,
                    empresas=len(empresas),
                    nomes=nomes,
                )
            )
        lines.sort(key=lambda line: (line.horario or "99:99", line.title.lower()))
        days.append(DayBucket(day=day, offset=_business_offset(anchor, day), lines=lines))
    return days


def _build_post_mortem(
    items: list[tuple[str, dict[str, Any], date | None, str, str]],
) -> list[PostMortemItem]:
    grouped: dict[tuple[str, date, str], dict[str, Any]] = {}
    for kind, row, prazo, title, _horario in items:
        if prazo is None:
            continue
        status = _status(row)
        delivered = _late_delivery(row, prazo, kind)
        if status != "ATRASADO" and delivered is None:
            continue
        key = (title, prazo, status)
        bucket = grouped.setdefault(
            key,
            {"empresas": set(), "delivered": delivered},
        )
        bucket["empresas"].add(_empresa_key(row))
        if delivered and (
            bucket["delivered"] is None or delivered > bucket["delivered"]
        ):
            bucket["delivered"] = delivered

    result: list[PostMortemItem] = []
    for (title, prazo, status), bucket in sorted(grouped.items(), key=lambda item: (item[0][1], item[0][0])):
        delivered = bucket["delivered"]
        if delivered:
            detalhe = f"Entregue em {delivered.strftime('%d/%m/%Y')} · prazo {prazo.strftime('%d/%m/%Y')}"
        else:
            detalhe = f"Atrasado · prazo {prazo.strftime('%d/%m/%Y')}"
        result.append(
            PostMortemItem(
                title=title,
                status=status,
                prazo=prazo,
                empresas=len(bucket["empresas"]),
                detalhe=detalhe,
            )
        )
    return result


FONT = "Barlow"
HIGHLIGHTS_PER_SLIDE = 8
RED = RGBColor(0xFF, 0x00, 0x00)


def _set_run(paragraph, text: str, *, size: int, bold: bool = False, color: RGBColor = INK, font: str = FONT) -> None:
    paragraph.clear()
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _textbox(slide, left, top, width, height, text: str, *, size: int, bold: bool = False, color: RGBColor = INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    try:
        frame._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(anchor, "t"))
    except Exception:
        pass
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    _set_run(paragraph, text, size=size, bold=bold, color=color)
    return box


def _rect(slide, left, top, width, height, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _fill(shape, color)
    return shape


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _period_code(competencia: date | None) -> str:
    if competencia is None:
        return "P"
    return f"P{competencia.month}"


def _period_badge(competencia: date | None) -> str:
    if competencia is None:
        return "Fechamento"
    return f"Fechamento P{competencia.month}/ {competencia.year % 100:02d}"


def _hora_label(horario: str) -> str:
    if not horario:
        return ""
    start = horario.split("–")[0].strip()
    if len(start) >= 5 and start[2] == ":":
        return f"{start[:2]}h{start[3:5]}"
    return start


def _flatten(days: list[DayBucket]) -> list[tuple[date, int, CalendarLine]]:
    rows: list[tuple[date, int, CalendarLine]] = []
    for bucket in days:
        for line in bucket.lines:
            rows.append((bucket.day, bucket.offset, line))
    return rows


def _runs_box(slide, left, top, width, height, entries: list[list[tuple[str, bool, RGBColor]]], *, size: int = 16):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    if not entries:
        entries = [[("Sem itens neste recorte.", False, MUTED)]]
    for index, parts in enumerate(entries):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(6)
        paragraph.clear()
        for text, bold, color in parts:
            run = paragraph.add_run()
            run.text = text
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            run.font.name = FONT
    return box


def _footer_bar(slide, lines: list[str]) -> None:
    _rect(slide, Inches(0), Inches(6.09), SLIDE_W, Inches(1.41), BRAND)
    text = "\n".join(line for line in lines if line) or "Market Call"
    _textbox(
        slide,
        Inches(0.42),
        Inches(6.22),
        Inches(12.2),
        Inches(1.12),
        text,
        size=16,
        bold=True,
        color=WHITE,
    )


def _section_title(slide, text: str) -> None:
    _textbox(slide, Inches(1.15), Inches(0.28), Inches(11.4), Inches(0.85), text, size=32, bold=True, color=INK)


def _highlight_parts(day: date, offset: int, line: CalendarLine) -> list[tuple[str, bool, RGBColor]]:
    hora = _hora_label(line.horario)
    title = line.title
    if line.empresas > 1:
        title = f"{title} ({line.empresas} empresas)"
    parts: list[tuple[str, bool, RGBColor]] = [
        (f"{day.strftime('%d/%m/%Y')} - ", True, INK),
        (_bd_label(offset), True, RED),
        (f" – {title}", False, INK),
    ]
    if hora:
        parts.append((f" – {hora}", True, INK))
    parts.append((";", False, INK))
    return parts


def _callout(day: date, offset: int, line: CalendarLine) -> str:
    hora = _hora_label(line.horario)
    suffix = f" – {hora}" if hora else ""
    return f"{_bd_label(offset)} – {line.title}{suffix};"


def _notes(rows: list[tuple[date, int, CalendarLine]]) -> list[list[tuple[str, bool, RGBColor]]]:
    notes: list[list[tuple[str, bool, RGBColor]]] = []
    seen: set[str] = set()
    for _day, _offset, line in rows:
        if len(line.nomes) < 2 or line.title in seen:
            continue
        seen.add(line.title)
        listed = ", ".join(line.nomes[:8])
        extra = "" if len(line.nomes) <= 8 else f" e mais {len(line.nomes) - 8}"
        notes.append([(f"Nota {len(notes) + 1}: {line.title} - {listed}{extra}", False, MUTED)])
        if len(notes) == 3:
            break
    return notes


def _add_cover(prs: Presentation, competencia: date | None, anchor: date) -> None:
    slide = _blank(prs)
    code = _period_code(competencia)
    badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.89),
        Inches(1.92),
        Inches(3.1),
        Inches(0.5),
    )
    _fill(badge, RGBColor(0xFF, 0xF1, 0xEC))
    _textbox(
        slide,
        Inches(0.89),
        Inches(2.0),
        Inches(3.1),
        Inches(0.36),
        _period_badge(competencia),
        size=16,
        bold=True,
        color=BRAND,
        align=PP_ALIGN.CENTER,
    )
    _textbox(slide, Inches(0.89), Inches(2.55), Inches(10), Inches(0.9), "Market Call . Brasil", size=40, bold=True, color=INK)

    agenda = (
        "Market Call BR - Introdução",
        f"Calendário – {code}: Highligths",
        "Gestão à Vista",
        "Prévias de Resultados",
        f"Post Mortem – {code}: Highlights",
        "Q&A",
    )
    width = Inches(12.1 / len(agenda))
    for index, label in enumerate(agenda):
        left = Inches(0.55) + width * index
        _textbox(slide, left, Inches(4.22), width, Inches(0.55), label, size=11, bold=True, color=BRAND)
        if index:
            _rect(slide, left, Inches(4.28), Inches(0.015), Inches(0.28), BRAND)

    _textbox(
        slide,
        Inches(0.55),
        Inches(4.82),
        Inches(12.2),
        Inches(0.4),
        "11H30  --------------------  11H35  ------------------------  11H40  --------------------  11H45  --------------------  11H50  -------------  11H55  -------  12H00",
        size=11,
        bold=True,
        color=BRAND,
    )
    _textbox(
        slide,
        Inches(10.4),
        Inches(6.45),
        Inches(2.4),
        Inches(0.4),
        anchor.strftime("%d/%m/%Y"),
        size=16,
        bold=True,
        color=BRAND,
        align=PP_ALIGN.RIGHT,
    )


def _add_calendar_slide(
    prs: Presentation,
    competencia: date | None,
    rows: list[tuple[date, int, CalendarLine]],
    total: int,
) -> None:
    slide = _blank(prs)
    code = _period_code(competencia)
    _section_title(slide, f"Calendário – {code}: Highligths")
    shown = rows[:HIGHLIGHTS_PER_SLIDE]
    entries = [_highlight_parts(day, offset, line) for day, offset, line in shown]
    if total > len(shown):
        entries.append([(f"E demais {total - len(shown)} atividades no cronograma filtrado.", False, MUTED)])
    if not entries:
        entries = [[("Nenhum destaque com prazo neste recorte.", False, MUTED)]]
    _runs_box(slide, Inches(1.04), Inches(1.35), Inches(11.8), Inches(3.55), entries, size=15)
    notes = _notes(shown)
    if notes:
        _runs_box(slide, Inches(1.04), Inches(4.95), Inches(11.5), Inches(1.0), notes, size=11)
    callouts = [_callout(day, offset, line) for day, offset, line in shown[:2]]
    _footer_bar(slide, callouts or ["Calendário"])


def _window_lines(rows: list[tuple[date, int, CalendarLine]], start: int | None, end: int | None) -> list[tuple[date, int, CalendarLine]]:
    picked = []
    for day, offset, line in rows:
        if start is not None and offset < start:
            continue
        if end is not None and offset > end:
            continue
        picked.append((day, offset, line))
    return picked


def _add_gestao(
    prs: Presentation,
    obrigacoes: list[dict[str, Any]],
    tarefas: list[dict[str, Any]],
) -> None:
    slide = _blank(prs)
    _section_title(slide, "Gestão à Vista")
    counts = {key: 0 for key in STATUS_LABELS}
    for row in [*obrigacoes, *tarefas]:
        status = _status(row)
        counts[status] = counts.get(status, 0) + 1
    late = counts.get("ATRASADO", 0)
    entries = [
        [
            ("Nesta competência serão monitorados os ", False, INK),
            ("prazos de entrega;", True, INK),
        ],
        [
            ("Obrigações: ", True, INK),
            (str(len(obrigacoes)), False, INK),
            ("    Tarefas: ", True, INK),
            (f"{len(tarefas)};", False, INK),
        ],
        [
            ("Pendente ", True, INK),
            (f"{counts.get('PENDENTE', 0)}; ", False, INK),
            ("Em andamento ", True, INK),
            (f"{counts.get('EM_ANDAMENTO', 0)}; ", False, INK),
            ("Em revisão ", True, INK),
            (f"{counts.get('EM_REVISAO', 0)};", False, INK),
        ],
        [
            ("Entregue ", True, INK),
            (f"{counts.get('ENTREGUE', 0)}; ", False, INK),
            ("Atrasado ", True, INK),
            (f"{late}.", False, INK),
        ],
        [
            (
                "O detalhe dos atrasos e das entregas fora do prazo está no post mortem."
                if late
                else "Nenhum item atrasado nesta visão filtrada.",
                False,
                INK,
            )
        ],
    ]
    _runs_box(slide, Inches(1.07), Inches(1.5), Inches(11.3), Inches(3.6), entries, size=16)
    _footer_bar(slide, ["Gestão à Vista"])


def _preview_line(label: str, rows: list[tuple[date, int, CalendarLine]]) -> list[tuple[str, bool, RGBColor]]:
    if not rows:
        return [(f"{label}: sem itens neste recorte;", False, INK)]
    titles = []
    for _day, _offset, line in rows:
        if line.title not in titles:
            titles.append(line.title)
    sample = "; ".join(titles[:3])
    extra = f" e mais {len(titles) - 3}" if len(titles) > 3 else ""
    return [
        (f"{label}: ", True, INK),
        (f"{len(rows)} atividades — {sample}{extra};", False, INK),
    ]


def _add_previas(prs: Presentation, rows: list[tuple[date, int, CalendarLine]]) -> None:
    slide = _blank(prs)
    _section_title(slide, "Prévias de Resultados")
    entries = [
        [("Expectativas das prévias de resultados:", True, INK)],
        _preview_line("Prévia 1 BD – 5", _window_lines(rows, None, -1)),
        _preview_line("Prévia 2 BD +2", _window_lines(rows, 0, 3)),
        _preview_line("Prévia 3 BD +4", _window_lines(rows, 4, None)),
    ]
    _runs_box(slide, Inches(1.07), Inches(1.5), Inches(11.3), Inches(3.6), entries, size=16)
    _footer_bar(slide, ["Prévias de Resultados"])


def _add_post_mortem(prs: Presentation, items: list[PostMortemItem], competencia: date | None) -> None:
    slide = _blank(prs)
    code = _period_code(competencia)
    _section_title(slide, f"Post Mortem – {code}: Highlights")
    shown = items[:8]
    entries: list[list[tuple[str, bool, RGBColor]]] = []
    for index, item in enumerate(shown, start=1):
        reason = "ENTREGA FORA DO PRAZO" if item.detalhe.startswith("Entregue") else "ATRASO"
        extra = f" – {item.empresas} EMPRESAS" if item.empresas > 1 else ""
        entries.append([(f"PM{index:02d} - {item.title.upper()}: {reason}{extra};", True, INK)])
    if len(items) > len(shown):
        entries.append([(f"E demais {len(items) - len(shown)} itens no cronograma filtrado.", False, MUTED)])
    if not entries:
        entries = [[("Nenhum item atrasado ou entregue depois do prazo.", False, MUTED)]]
    _runs_box(slide, Inches(1.07), Inches(1.45), Inches(11.3), Inches(3.8), entries, size=15)
    _footer_bar(slide, [f"Post Mortem – {code}"])


def _add_qa(prs: Presentation) -> None:
    slide = _blank(prs)
    _textbox(slide, Inches(1.15), Inches(2.05), Inches(10), Inches(1.1), "Q&A", size=40, bold=True, color=INK)
    _footer_bar(slide, ["11H55 – 12H00", "Q&A"])


def _add_closing(prs: Presentation) -> None:
    slide = _blank(prs)
    _textbox(slide, Inches(1.15), Inches(2.05), Inches(10), Inches(1.1), "OBRIGADO", size=40, bold=True, color=INK)
    _footer_bar(slide, ["Market Call", "De Segunda à Sexta-Feira | 11H00 | Gestores CSC"])


def build_market_call_pptx(
    obrigacoes: list[dict[str, Any]],
    competencia: date | None = None,
    tarefas: list[dict[str, Any]] | None = None,
) -> bytes:
    tarefas = tarefas or []
    anchor = _resolve_anchor(competencia, obrigacoes + tarefas)
    items = _collect_items(obrigacoes, tarefas)
    days = _build_days(items, anchor)
    rows = _flatten(days)
    before = [row for row in rows if row[1] <= 0]
    after = [row for row in rows if row[1] > 0]
    post_mortem = _build_post_mortem(items)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    _add_cover(prs, competencia, anchor)
    _add_calendar_slide(prs, competencia, before, len(before))
    _add_calendar_slide(prs, competencia, after, len(after))
    _add_gestao(prs, obrigacoes, tarefas)
    _add_previas(prs, rows)
    _add_post_mortem(prs, post_mortem, competencia)
    _add_qa(prs)
    _add_closing(prs)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
