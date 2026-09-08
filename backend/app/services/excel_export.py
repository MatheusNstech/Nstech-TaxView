from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BRAND = "FF3D03"
INK = "1F2937"
MUTED = "6B7280"
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, color=INK, size=16)
SUB_FONT = Font(color=MUTED, size=11)
ZEBRA = PatternFill("solid", fgColor="FFF7F5")
THIN = Border(
    left=Side(style="thin", color="F3E8E4"),
    right=Side(style="thin", color="F3E8E4"),
    top=Side(style="thin", color="F3E8E4"),
    bottom=Side(style="thin", color="F3E8E4"),
)

STATUS_LABELS = {
    "PENDENTE": "Pendente",
    "EM_ANDAMENTO": "Em andamento",
    "EM_REVISAO": "Em revisão",
    "ENTREGUE": "Entregue",
    "ATRASADO": "Atrasado",
}
STATUS_FILL = {
    "PENDENTE": PatternFill("solid", fgColor="F3F4F6"),
    "EM_ANDAMENTO": PatternFill("solid", fgColor="DBEAFE"),
    "EM_REVISAO": PatternFill("solid", fgColor="FEF3C7"),
    "ENTREGUE": PatternFill("solid", fgColor="D1FAE5"),
    "ATRASADO": PatternFill("solid", fgColor="FEE2E2"),
}
STATUS_FONT = {
    "PENDENTE": Font(color="374151", bold=True, size=10),
    "EM_ANDAMENTO": Font(color="1D4ED8", bold=True, size=10),
    "EM_REVISAO": Font(color="B45309", bold=True, size=10),
    "ENTREGUE": Font(color="047857", bold=True, size=10),
    "ATRASADO": Font(color="B91C1C", bold=True, size=10),
}
URGENCIA_LABELS = {
    "ok": "No prazo",
    "urgente": "Urgente",
    "atrasado": "Atrasado",
    "neutro": "Sem prazo",
}

HEADERS = [
    "Tipo",
    "Empresa",
    "CNPJ",
    "BU",
    "Atividade / Título",
    "Responsável",
    "Solicitante",
    "Competência",
    "Prazo legal",
    "Prazo",
    "Horário",
    "Data entrega",
    "Status",
    "Urgência",
]
LAST_COL = get_column_letter(len(HEADERS))
TIPO_FILL = {
    "Obrigação": PatternFill("solid", fgColor="FFF1EC"),
    "Tarefa": PatternFill("solid", fgColor="EEF2FF"),
}
TIPO_FONT = {
    "Obrigação": Font(color=BRAND, bold=True, size=10),
    "Tarefa": Font(color="3730A3", bold=True, size=10),
}


def _as_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value)[:10]
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def _month_label(value: date | None) -> str:
    if value is None:
        return "Todas as competências"
    months = (
        "janeiro",
        "fevereiro",
        "março",
        "abril",
        "maio",
        "junho",
        "julho",
        "agosto",
        "setembro",
        "outubro",
        "novembro",
        "dezembro",
    )
    return f"{months[value.month - 1]} de {value.year}"


def _hora(value: Any) -> str:
    if value is None or value == "":
        return ""
    text = str(value).strip()
    return text[:5] if text else ""


def _horario(inicio: Any, fim: Any) -> str:
    start = _hora(inicio)
    end = _hora(fim)
    if start and end:
        return f"{start}–{end}"
    return start or end


def _row_values(kind: str, row: dict[str, Any]) -> tuple[list[Any], str]:
    emp = row.get("empresa") or {}
    resp = row.get("responsavel") or {}
    status = str(row.get("status") or "PENDENTE")
    urgencia = str(row.get("urgencia") or "")
    if kind == "Tarefa":
        entrega = _as_date(row.get("entregue_em"))
        item = row.get("titulo") or ""
        prazo = _as_date(row.get("prazo"))
        values: list[Any] = [
            "Tarefa",
            emp.get("razao_social") or "",
            emp.get("cnpj") or "",
            emp.get("bu") or "",
            item,
            resp.get("nome") or "",
            row.get("solicitante_nome") or "",
            _as_date(row.get("competencia")),
            None,
            prazo,
            _horario(row.get("hora_inicio"), row.get("hora_fim")),
            entrega,
            STATUS_LABELS.get(status, status),
            URGENCIA_LABELS.get(urgencia, urgencia),
        ]
        return values, status
    atv = row.get("atividade") or {}
    values = [
        "Obrigação",
        emp.get("razao_social") or "",
        emp.get("cnpj") or "",
        emp.get("bu") or "",
        atv.get("nome") or "",
        resp.get("nome") or "",
        "",
        _as_date(row.get("competencia")),
        _as_date(row.get("prazo_legal")),
        _as_date(row.get("prazo_fiscal")),
        "",
        _as_date(row.get("data_entrega")),
        STATUS_LABELS.get(status, status),
        URGENCIA_LABELS.get(urgencia, urgencia),
    ]
    return values, status


def build_obrigacoes_xlsx(
    rows: list[dict[str, Any]],
    competencia: date | None = None,
    tarefas: list[dict[str, Any]] | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = "MsCronograma"
    ws.oddFooter.right.text = "Página &P de &N"

    tarefas = tarefas or []
    combined: list[tuple[str, dict[str, Any]]] = [
        ("Obrigação", row) for row in rows
    ] + [("Tarefa", row) for row in tarefas]
    counts: dict[str, int] = {key: 0 for key in STATUS_LABELS}
    for _, row in combined:
        status = str(row.get("status") or "PENDENTE")
        counts[status] = counts.get(status, 0) + 1

    ws.merge_cells(f"A1:{LAST_COL}1")
    ws["A1"] = "Cronograma de obrigações e tarefas"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells(f"A2:{LAST_COL}2")
    ws["A2"] = (
        f"{_month_label(competencia)}  ·  {len(rows)} obrigações  ·  "
        f"{len(tarefas)} tarefas  ·  "
        f"Pendente {counts.get('PENDENTE', 0)}  ·  "
        f"Em andamento {counts.get('EM_ANDAMENTO', 0)}  ·  "
        f"Em revisão {counts.get('EM_REVISAO', 0)}  ·  "
        f"Entregue {counts.get('ENTREGUE', 0)}  ·  "
        f"Atrasado {counts.get('ATRASADO', 0)}"
    )
    ws["A2"].font = SUB_FONT
    ws.row_dimensions[2].height = 18

    header_row = 4
    for col, name in enumerate(HEADERS, start=1):
        cell = ws.cell(header_row, col, name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = THIN
    ws.row_dimensions[header_row].height = 22
    ws.auto_filter.ref = f"A{header_row}:{LAST_COL}{header_row}"
    ws.freeze_panes = "A5"
    ws.print_title_rows = "1:4"

    body_font = Font(color=INK, size=10)
    date_cols = {8, 9, 10, 12}
    status_col = HEADERS.index("Status") + 1
    tipo_col = 1
    cnpj_col = HEADERS.index("CNPJ") + 1
    for index, (kind, row) in enumerate(combined):
        excel_row = header_row + 1 + index
        values, status = _row_values(kind, row)
        zebra = index % 2 == 1
        for col, value in enumerate(values, start=1):
            cell = ws.cell(excel_row, col, value)
            cell.font = body_font
            cell.border = THIN
            cell.alignment = Alignment(vertical="center")
            if zebra and col not in {tipo_col, status_col}:
                cell.fill = ZEBRA
            if col in date_cols and isinstance(value, date):
                cell.number_format = "DD/MM/YYYY"
            if col == cnpj_col:
                cell.number_format = "@"
            if col == tipo_col:
                cell.fill = TIPO_FILL[kind]
                cell.font = TIPO_FONT[kind]
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if col == status_col:
                cell.fill = STATUS_FILL.get(status, STATUS_FILL["PENDENTE"])
                cell.font = STATUS_FONT.get(status, STATUS_FONT["PENDENTE"])
                cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[excel_row].height = 20

    widths = [14, 36, 20, 16, 34, 24, 24, 16, 15, 15, 16, 15, 16, 14]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    if combined:
        ws.auto_filter.ref = f"A{header_row}:{LAST_COL}{header_row + len(combined)}"

    ws.oddHeader.center.text = _month_label(competencia)
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.tabColor = BRAND

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
